"""Trich text tu docx/pdf/xlsx/pptx/rtf/txt/doc de chung cat.
Dung: python tools/extract_text.py <file> [--max N]   (N = so ky tu toi da, mac dinh 40000)
"""
import sys, io, re, zipfile
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")


def docx_text(path):
    try:
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        if parts:
            return "\n".join(parts)
    except Exception:
        pass
    # fallback: loi text tho tu XML trong zip (docx loi hoac text nam trong SmartArt/diagram)
    out = []
    with zipfile.ZipFile(path) as z:
        for n in sorted(z.namelist()):
            if n.endswith(".xml") and (n.startswith("word/document") or n.startswith("word/diagrams")
                                       or n.startswith("word/header") or n.startswith("word/footer")):
                xml = z.read(n).decode("utf-8", errors="replace")
                txts = re.findall(r"<(?:w|a):t[^>]*>([^<]*)</(?:w|a):t>", xml)
                if txts:
                    out.append(f"[{n}] " + " ".join(t for t in txts if t.strip()))
    return "\n".join(out)


def pdf_text(path):
    from pypdf import PdfReader
    r = PdfReader(path)
    return "\n".join((p.extract_text() or "") for p in r.pages)


def xlsx_text(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"### Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None and str(v).strip()]
            if vals:
                out.append(" | ".join(vals))
        if len("\n".join(out)) > 60000:
            out.append("...[sheet cat bot]")
            break
    return "\n".join(out)


def pptx_text(path):
    try:
        from pptx import Presentation
        pr = Presentation(path)
        out = []
        for i, sl in enumerate(pr.slides, 1):
            for sh in sl.shapes:
                if sh.has_text_frame:
                    t = sh.text_frame.text.strip()
                    if t:
                        out.append(f"[slide {i}] {t}")
        return "\n".join(out)
    except ImportError:
        # fallback: lấy text thô từ xml trong zip
        out = []
        with zipfile.ZipFile(path) as z:
            for n in sorted(z.namelist()):
                if n.startswith("ppt/slides/slide") and n.endswith(".xml"):
                    xml = z.read(n).decode("utf-8", errors="replace")
                    txts = re.findall(r"<a:t>([^<]*)</a:t>", xml)
                    if txts:
                        out.append(f"[{n}] " + " ".join(txts))
        return "\n".join(out)


def rtf_text(path):
    raw = open(path, "rb").read().decode("latin-1", errors="replace")
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", raw)
    raw = re.sub(r"[{}]", "", raw)
    return re.sub(r"\s+", " ", raw).strip()


def old_doc_text(path):
    # .doc nhị phân: lôi chuỗi UTF-16LE có dấu tiếng Việt
    raw = open(path, "rb").read()
    try:
        txt = raw.decode("utf-16-le", errors="ignore")
    except Exception:
        txt = raw.decode("latin-1", errors="replace")
    chunks = re.findall(r"[\w\sÀ-ỹ.,;:!?%()/\-\n]{8,}", txt)
    return "\n".join(c.strip() for c in chunks if len(c.strip()) > 8)


def extract(path):
    low = path.lower()
    if low.endswith(".docx"):
        return docx_text(path)
    if low.endswith(".pdf"):
        return pdf_text(path)
    if low.endswith(".xlsx"):
        return xlsx_text(path)
    if low.endswith(".pptx"):
        return pptx_text(path)
    if low.endswith(".rtf"):
        return rtf_text(path)
    if low.endswith((".txt", ".md", ".csv")):
        return open(path, encoding="utf-8", errors="replace").read()
    if low.endswith(".doc"):
        return old_doc_text(path)
    return f"[dinh dang chua ho tro: {path}]"


if __name__ == "__main__":
    path = sys.argv[1]
    mx = 40000
    if "--max" in sys.argv:
        mx = int(sys.argv[sys.argv.index("--max") + 1])
    try:
        print(extract(path)[:mx])
    except Exception as e:
        print(f"[LOI TRICH: {type(e).__name__}: {e}]")
